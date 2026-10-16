"""Independently reprice the six MO2610 report scenarios with Python BSM.

Inputs: the report DOCX (legs, entry premiums, scenarios, reported PnL), the
original GB18030 option-chain CSV (per-strike IV), and explicit r/q assumptions.
The report itself does not record r or q. Run from the repository root:

    python3 reports/07_mo2610_python_bsm_reconcile.py
"""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from datetime import date, timedelta
from math import erf, exp, log, sqrt
from pathlib import Path
import re

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "MO2610_期权情景分析报告.docx"
CHAIN = ROOT / "options_strategy_workbench/report.csv"
MULTIPLIER = 100


@dataclass(frozen=True)
class Leg:
    side: int
    option_type: str
    strike: int
    premium: float
    iv: float


def number(value: str) -> float:
    return float(re.sub(r"\s", "", value).replace("％", "%").rstrip("%"))


def normal_cdf(value: float) -> float:
    return 0.5 * (1 + erf(value / sqrt(2)))


def bsm_value(spot: float, strike: float, sigma: float, years: float,
              rate: float, dividend: float, option_type: str) -> float:
    if years <= 0:
        return max(spot - strike, 0) if option_type == "CALL" else max(strike - spot, 0)
    d1 = (log(spot / strike) + (rate - dividend + sigma * sigma / 2) * years) / (sigma * sqrt(years))
    d2 = d1 - sigma * sqrt(years)
    if option_type == "CALL":
        return spot * exp(-dividend * years) * normal_cdf(d1) - strike * exp(-rate * years) * normal_cdf(d2)
    return strike * exp(-rate * years) * normal_cdf(-d2) - spot * exp(-dividend * years) * normal_cdf(-d1)


def read_chain_iv() -> dict[tuple[str, int], float]:
    result = {}
    with CHAIN.open(encoding="gb18030", newline="") as handle:
        for row in list(csv.reader(handle))[1:]:
            strike = int(re.search(r"\d+", row[9]).group())
            result["CALL", strike] = number(row[7]) / 100
            result["PUT", strike] = number(row[11]) / 100
    return result


def read_report():
    doc = Document(REPORT)
    iv_by_leg = read_chain_iv()
    overview = doc.tables[0]
    expiry = date.fromisoformat(overview.cell(0, 3).text.strip())
    spot = number(overview.cell(1, 1).text.split("］")[0].split("［")[-1])
    month, day = map(int, re.search(r"［\s*(\d{1,2})\.(\d{1,2})\s+\d{1,2}:\d{2}", overview.cell(1, 1).text).groups())
    valuation = date(expiry.year, month, day)

    strategies = {}
    for name, table in (("Bear Call", doc.tables[6]), ("Iron Condor", doc.tables[9])):
        legs = []
        for row in table.rows[1:]:
            cells = [c.text for c in row.cells]
            side = -1 if re.sub(r"\s", "", cells[1]).lower() == "short" else 1
            option_type = re.sub(r"\s", "", cells[2]).upper()
            strike = int(number(cells[3]))
            legs.append(Leg(side, option_type, strike, number(cells[5]), iv_by_leg[option_type, strike]))
        strategies[name] = legs

    scenarios = []
    for row, result in zip(doc.tables[13].rows[1:], doc.tables[14].rows[1:7], strict=True):
        cells = [c.text for c in row.cells]
        reported = [c.text for c in result.cells]
        scenarios.append((cells[0].strip(), number(cells[1]) / 100, number(cells[2]) / 100,
                          int(number(cells[3].split("天")[0])), number(cells[5]) / 100,
                          number(reported[1]), number(reported[2])))
    return spot, valuation, expiry, strategies, scenarios, iv_by_leg


def pnl(legs: list[Leg], spot: float, years: float, iv_shift: float,
        rate: float, dividend: float) -> tuple[float, list[float]]:
    marks = [bsm_value(spot, leg.strike, leg.iv + iv_shift, years,
                       rate, dividend, leg.option_type) for leg in legs]
    value = sum(leg.side * (mark - leg.premium) * MULTIPLIER for leg, mark in zip(legs, marks, strict=True))
    return value, marks


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rate", type=float, default=1.65, help="annual risk-free rate in percent; absent from the DOCX")
    parser.add_argument("--dividend", type=float, default=1.1, help="annual MO dividend yield in percent; absent from the DOCX")
    args = parser.parse_args()
    spot, valuation, expiry, strategies, scenarios, iv_by_leg = read_report()
    rate, dividend = args.rate / 100, args.dividend / 100
    print(f"Spot {spot:.2f}; valuation {valuation}; expiry {expiry}; T0={(expiry-valuation).days}/365; "
          f"r={args.rate}%; q={args.dividend}%; multiplier={MULTIPLIER}")
    for name, legs in strategies.items():
        print(name, [("SHORT" if leg.side < 0 else "LONG", leg.option_type, leg.strike,
                       leg.premium, round(leg.iv * 100, 2)) for leg in legs])
    print("scenario,report_bear,python_bear,difference,report_iron,python_iron,difference,saved_C6900_C7100")
    saved_bear = [Leg(-1, "CALL", 6900, 80.8, iv_by_leg["CALL", 6900]),
                  Leg(1, "CALL", 7100, 16.2, iv_by_leg["CALL", 7100])]
    weighted = [0.0, 0.0]
    for name, price_move, iv_move, days_forward, weight, report_bear, report_iron in scenarios:
        remaining = max((expiry - valuation - timedelta(days=days_forward)).days, 0) / 365
        scenario_spot = spot * (1 + price_move)
        bear, _ = pnl(strategies["Bear Call"], scenario_spot, remaining, iv_move, rate, dividend)
        iron, _ = pnl(strategies["Iron Condor"], scenario_spot, remaining, iv_move, rate, dividend)
        saved, _ = pnl(saved_bear, scenario_spot, remaining, iv_move, rate, dividend)
        weighted[0] += weight * bear
        weighted[1] += weight * iron
        print(f"{name},{report_bear:.2f},{bear:.2f},{bear-report_bear:+.2f},"
              f"{report_iron:.2f},{iron:.2f},{iron-report_iron:+.2f},{saved:.2f}")
    print(f"weighted_python_bear={weighted[0]:.2f}; weighted_python_iron={weighted[1]:.2f}")


if __name__ == "__main__":
    main()
